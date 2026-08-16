"""Registro de metricas de simulacao numa planilha .xlsx acumulada (1 linha por run).

Pensado pra comparar simulacoes entre si (ex.: `explicit` vs `semi_lagrangian`,
diferentes `dt`/malha/Re) numa tabela unica, nao pra analise por-iteracao
dentro de um unico run -- ver `scripts/run_simulation.py` sobre quais
metricas sao agregadas (media/max) antes de virar uma linha aqui.
"""

from pathlib import Path

from openpyxl import Workbook, load_workbook

COLUNAS = [
    "timestamp", "mesh", "advection", "sl_boundary", "element", "dt", "reynolds", "iterations",
    "npoints", "ne", "device",
    "tempo_total_s", "tempo_assembly_s", "tempo_medio_por_iter_s",
    "bicg_iters_media", "bicg_iters_max",
    "bicg_residual_media", "bicg_residual_max",
    "passos_nao_convergidos",
    "vel_l2_media", "vel_l2_max", "vel_linf_max",
    "divergencia_l2_media", "divergencia_l2_max",
    "pressao_min", "pressao_max", "pressao_media",
]


def _migrar_cabecalho_se_preciso(ws):
    """Reescreve a planilha com o cabecalho atual (`COLUNAS`) se o arquivo
    foi criado com uma versao antiga (ex.: uma coluna nova, como `element`,
    foi acrescentada ao meio de `COLUNAS` depois que o arquivo ja existia).

    Sem isso, linhas novas (escritas na ordem de `COLUNAS` atual) ficam
    desalinhadas do cabecalho antigo salvo na planilha -- os dados ficam
    certos na celula, mas o nome da coluna no cabecalho nao bate mais a
    partir do ponto de insercao. Linhas antigas ganham celula vazia nas
    colunas que nao existiam quando foram escritas (mesma regra de `dados`
    ausente em `salvar_resumo`).
    """
    cabecalho_salvo = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    if cabecalho_salvo == COLUNAS:
        return

    linhas_antigas = [dict(zip(cabecalho_salvo, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]

    ws.delete_rows(1, ws.max_row)
    ws.append(COLUNAS)
    for linha in linhas_antigas:
        ws.append([linha.get(coluna, "") for coluna in COLUNAS])


def salvar_resumo(path: str, dados: dict):
    """Acrescenta uma linha de resumo em `path` (.xlsx), criando arquivo/cabecalho se preciso.

    `dados` mapeia (um subconjunto de) `COLUNAS` para valores dessa
    simulacao; chaves de `COLUNAS` ausentes em `dados` viram celula vazia
    (nao quebra se uma metrica nao foi calculada) e chaves extras em
    `dados` sao ignoradas (nao aparecem na planilha).
    """
    caminho = Path(path)

    if caminho.exists():
        wb = load_workbook(caminho)
        ws = wb.active
        _migrar_cabecalho_se_preciso(ws)
    else:
        caminho.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = "Resumo"
        ws.append(COLUNAS)

    ws.append([dados.get(coluna, "") for coluna in COLUNAS])
    wb.save(caminho)
