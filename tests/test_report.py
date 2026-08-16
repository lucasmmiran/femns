from openpyxl import Workbook, load_workbook

from femns.report import COLUNAS, salvar_resumo


def test_salvar_resumo_cria_arquivo_com_cabecalho(tmp_path):
    path = tmp_path / "benchmarks.xlsx"
    salvar_resumo(path, {"mesh": "meshes/poiseuille.msh", "advection": "explicit", "dt": 0.001})

    wb = load_workbook(path)
    ws = wb.active
    linhas = list(ws.iter_rows(values_only=True))

    assert linhas[0] == tuple(COLUNAS)
    assert len(linhas) == 2

    linha = dict(zip(COLUNAS, linhas[1]))
    assert linha["mesh"] == "meshes/poiseuille.msh"
    assert linha["advection"] == "explicit"
    assert linha["dt"] == 0.001


def test_salvar_resumo_chave_ausente_vira_celula_vazia(tmp_path):
    path = tmp_path / "benchmarks.xlsx"
    salvar_resumo(path, {"mesh": "meshes/poiseuille.msh"})

    wb = load_workbook(path)
    ws = wb.active
    linha = dict(zip(COLUNAS, next(ws.iter_rows(min_row=2, values_only=True))))

    # openpyxl le uma celula de string vazia de volta como None
    assert linha["advection"] is None


def test_salvar_resumo_chave_extra_e_ignorada(tmp_path):
    path = tmp_path / "benchmarks.xlsx"
    salvar_resumo(path, {"mesh": "meshes/poiseuille.msh", "chave_que_nao_existe": 42})

    wb = load_workbook(path)
    ws = wb.active
    linha = dict(zip(COLUNAS, next(ws.iter_rows(min_row=2, values_only=True))))

    assert "chave_que_nao_existe" not in linha
    assert linha["mesh"] == "meshes/poiseuille.msh"


def test_salvar_resumo_acumula_uma_linha_por_chamada(tmp_path):
    path = tmp_path / "benchmarks.xlsx"
    salvar_resumo(path, {"mesh": "a", "advection": "explicit"})
    salvar_resumo(path, {"mesh": "a", "advection": "semi_lagrangian"})
    salvar_resumo(path, {"mesh": "b", "advection": "explicit"})

    wb = load_workbook(path)
    ws = wb.active
    linhas = list(ws.iter_rows(min_row=2, values_only=True))

    assert len(linhas) == 3
    advections = [dict(zip(COLUNAS, linha))["advection"] for linha in linhas]
    assert advections == ["explicit", "semi_lagrangian", "explicit"]


def test_salvar_resumo_cria_diretorio_pai_se_nao_existir(tmp_path):
    path = tmp_path / "subdir" / "benchmarks.xlsx"
    salvar_resumo(path, {"mesh": "a"})

    assert path.exists()


def test_salvar_resumo_migra_arquivo_com_cabecalho_antigo(tmp_path):
    """Simula um .xlsx criado com uma versao antiga de COLUNAS (sem colunas
    que so foram acrescentadas depois, ex. `element`/`sl_boundary`) --
    salvar_resumo deve realinhar o cabecalho pra versao atual em vez de
    desalinhar as linhas novas do cabecalho antigo salvo no arquivo.

    A linha antiga e' escrita a partir de `valores_antigos`, nao por
    posicao fixa, pra o teste continuar valendo quando outra coluna nova
    for acrescentada ao meio de COLUNAS.
    """
    path = tmp_path / "benchmarks.xlsx"

    novas = {"element", "sl_boundary"}
    cabecalho_antigo = [c for c in COLUNAS if c not in novas]
    valores_antigos = {
        "timestamp": "2026-01-01T00:00:00", "mesh": "meshes/poiseuille.msh",
        "advection": "explicit", "dt": 0.001, "reynolds": 1, "iterations": 1000,
    }
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.append(cabecalho_antigo)
    ws.append([valores_antigos.get(c) for c in cabecalho_antigo])
    wb.save(path)

    salvar_resumo(path, {"mesh": "meshes/degrau.msh", "advection": "explicit", "element": "tri6"})

    wb2 = load_workbook(path)
    ws2 = wb2.active
    linhas = list(ws2.iter_rows(values_only=True))

    assert linhas[0] == tuple(COLUNAS)  # cabecalho realinhado

    linha_antiga = dict(zip(COLUNAS, linhas[1]))
    assert linha_antiga["mesh"] == "meshes/poiseuille.msh"
    assert linha_antiga["dt"] == 0.001
    assert linha_antiga["element"] is None  # nao existia quando essa linha foi escrita

    linha_nova = dict(zip(COLUNAS, linhas[2]))
    assert linha_nova["mesh"] == "meshes/degrau.msh"
    assert linha_nova["element"] == "tri6"
