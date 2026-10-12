#!/usr/bin/env python3
"""CLI: compara os solvers lineares (BiCGSTAB x MINRES, com e sem pre-condicionador) lado a lado.

Para cada caso (um config do projeto) e cada variante de solver, roda
`scripts/run_simulation.py` com a **malha interior movel ligada**
(`mesh_motion.enabled: true`) e com a planilha de benchmark
(`femns.report`), e no fim monta uma tabela com:

- robustez: passos que nao atingiram a tolerancia (residuo verdadeiro) e o
  maior erro contra a solucao direta (velocidade e pressao), checado a cada
  `--verificar-direto` passos;
- tempo por passo (medio, maximo, minimo, sem a checagem direta) e tempo total;
- iteracoes do solver linear por passo (media, maxima, minima).

A tabela vai para `<saida>/relatorio.md`; as linhas completas, para
`<saida>/resultados.xlsx`. Plano 12/10 (`docs/planos/plano1210.md`).

Exemplo::

    python scripts/benchmark_solvers.py --iteracoes 100 --verificar-direto 20
"""

import argparse
import copy
import os
import subprocess
import sys
from datetime import datetime

import yaml
from openpyxl import load_workbook

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# (rotulo, opcoes de simulation.solver)
VARIANTES = [
    ("B0 bicgstab", {"metodo": "bicgstab", "precondicionador": "nenhum"}),
    ("B2 bicgstab + bloco-triangular (chebyshev)", {"metodo": "bicgstab", "precondicionador": "bloco_triangular",
                                                     "subsolver": "chebyshev"}),
    ("B2 bicgstab + bloco-triangular (lu)", {"metodo": "bicgstab", "precondicionador": "bloco_triangular",
                                              "subsolver": "lu"}),
    ("B2 bicgstab + bloco-triangular (fsai)", {"metodo": "bicgstab", "precondicionador": "bloco_triangular",
                                                "subsolver": "fsai",
                                                "chebyshev": {"velocidade": 5, "massa": 5, "laplaciano": 10}}),
    ("M0 minres", {"metodo": "minres", "precondicionador": "nenhum"}),
    ("M1 minres + bloco-diagonal (chebyshev)", {"metodo": "minres", "precondicionador": "bloco_diagonal",
                                                 "subsolver": "chebyshev"}),
    ("M1 minres + bloco-diagonal (lu)", {"metodo": "minres", "precondicionador": "bloco_diagonal",
                                          "subsolver": "lu"}),
    ("M1 minres + bloco-diagonal (fsai)", {"metodo": "minres", "precondicionador": "bloco_diagonal",
                                            "subsolver": "fsai",
                                            "chebyshev": {"velocidade": 5, "massa": 5, "laplaciano": 10}}),
]


def _slug(rotulo: str) -> str:
    """Nome de pasta da variante: o codigo (B2, M1...) mais o sub-solver, se houver."""
    codigo = rotulo.split()[0]
    sub = rotulo[rotulo.find("(") + 1:rotulo.find(")")] if "(" in rotulo else ""
    return f"{codigo}_{sub}" if sub else codigo


def parse_args():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--casos", nargs="+", default=["configs/poiseuille.yaml", "configs/degrau.yaml"])
    p.add_argument("--iteracoes", type=int, default=100, help="passos de tempo por rodada")
    p.add_argument("--verificar-direto", type=int, default=20, help="checagem contra o solver direto a cada N passos")
    p.add_argument("--variantes", nargs="*", default=None, help="prefixos dos rotulos a rodar (padrao: todas)")
    p.add_argument("--saida", default="solucoes/benchmark_solvers")
    p.add_argument("--tol", type=float, default=None, help="sobrescreve simulation.solver.tol do caso")
    return p.parse_args()


def config_da_rodada(base: dict, opcoes: dict, args, saida_caso: str, xlsx: str) -> dict:
    cfg = copy.deepcopy(base)
    cfg["output_dir"] = saida_caso
    cfg["benchmark_xlsx"] = xlsx
    cfg["simulation"]["iterations"] = args.iteracoes
    cfg["simulation"]["vtk_interval"] = args.iteracoes  # so o ultimo frame
    solver = dict(cfg["simulation"].get("solver") or {})
    solver.update(opcoes)
    solver["verificar_direto"] = args.verificar_direto
    if args.tol is not None:
        solver["tol"] = args.tol
    cfg["simulation"]["solver"] = solver
    cfg.setdefault("mesh_motion", {})["enabled"] = True  # sempre com malha interior movel
    return cfg


def main():
    args = parse_args()
    os.makedirs(args.saida, exist_ok=True)
    xlsx = os.path.join(args.saida, "resultados.xlsx")
    inicio = datetime.now().isoformat(timespec="seconds")
    variantes = [v for v in VARIANTES if not args.variantes or any(v[0].startswith(p) for p in args.variantes)]

    rodadas = []
    for caso in args.casos:
        with open(caso) as f:
            base = yaml.safe_load(f)
        nome_caso = os.path.splitext(os.path.basename(caso))[0]
        for rotulo, opcoes in variantes:
            pasta = os.path.join(args.saida, nome_caso, _slug(rotulo))
            cfg = config_da_rodada(base, opcoes, args, pasta, xlsx)
            caminho_cfg = os.path.join(args.saida, "configs", f"{nome_caso}_{_slug(rotulo)}.yaml")
            os.makedirs(os.path.dirname(caminho_cfg), exist_ok=True)
            with open(caminho_cfg, "w") as f:
                yaml.safe_dump(cfg, f, sort_keys=False)
            print(f"[{datetime.now():%H:%M:%S}] {nome_caso} :: {rotulo}", flush=True)
            log = os.path.join(args.saida, "configs", f"{nome_caso}_{_slug(rotulo)}.log")
            with open(log, "w") as f:
                ret = subprocess.run([sys.executable, os.path.join(RAIZ, "scripts", "run_simulation.py"),
                                      "--config", caminho_cfg], stdout=f, stderr=subprocess.STDOUT, cwd=RAIZ)
            rodadas.append((nome_caso, rotulo, ret.returncode, log))
            if ret.returncode != 0:
                print(f"    FALHOU (codigo {ret.returncode}), ver {log}", flush=True)

    escrever_relatorio(args, xlsx, inicio, rodadas)


def escrever_relatorio(args, xlsx, inicio, rodadas):
    linhas = []
    if os.path.exists(xlsx):
        ws = load_workbook(xlsx).active
        todas = list(ws.iter_rows(values_only=True))
        cab = todas[0]
        linhas = [dict(zip(cab, r)) for r in todas[1:] if str(dict(zip(cab, r)).get("timestamp", "")) >= inicio]

    def fmt(v, f="{:.3g}"):
        return "" if v in (None, "") else f.format(v)

    tabela = ["| caso | variante | passos sem convergir | erro direto v | erro direto p | t/passo medio (s) "
              "| t/passo max | t/passo min | t total (s) | iters media | iters max | iters min |",
              "|---|---|---|---|---|---|---|---|---|---|---|---|"]
    por_rodada = iter(linhas)
    for nome_caso, rotulo, codigo, log in rodadas:
        if codigo != 0:
            tabela.append(f"| {nome_caso} | {rotulo} | **falhou** (ver `{os.path.relpath(log, args.saida)}`) |" + " |" * 9)
            continue
        r = next(por_rodada)
        tabela.append("| " + " | ".join([
            nome_caso, rotulo, str(r["passos_nao_convergidos"]),
            fmt(r["erro_direto_v_max"], "{:.1e}"), fmt(r["erro_direto_p_max"], "{:.1e}"),
            fmt(r["tempo_passo_media_s"], "{:.3f}"), fmt(r["tempo_passo_max_s"], "{:.3f}"),
            fmt(r["tempo_passo_min_s"], "{:.3f}"), fmt(r["tempo_total_s"], "{:.1f}"),
            fmt(r["bicg_iters_media"], "{:.1f}"), str(r["bicg_iters_max"]), str(r["bicg_iters_min"]),
        ]) + " |")

    texto = (f"# Benchmark de solvers lineares ({inicio})\n\n"
             f"Malha interior movel ligada; {args.iteracoes} passos por rodada; checagem contra o solver "
             f"direto a cada {args.verificar_direto} passos (fora do tempo por passo); "
             f"tolerancia {args.tol if args.tol is not None else 'do config'}.\n\n"
             + "\n".join(tabela) + "\n")
    with open(os.path.join(args.saida, "relatorio.md"), "w") as f:
        f.write(texto)
    print("\n" + texto)


if __name__ == "__main__":
    main()
